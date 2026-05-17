import os
import json
from datetime import datetime

from fastapi import HTTPException, status
from sqlalchemy.orm import Session

from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

from app.models.report import Report
from app.models.forecast import ForecastHistory
from app.models.dataset import Dataset
from app.services.notification_service import create_notification


REPORT_DIR = "reports"


def parse_json(value, default=None):
    if not value:
        return default

    try:
        return json.loads(value)
    except Exception:
        return value


def get_forecast_or_404(db: Session, forecast_id: int, user_id: int, role: str):
    forecast = db.query(ForecastHistory).filter(
        ForecastHistory.id == forecast_id
    ).first()

    if not forecast:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail="Forecast not found"
        )

    if role != "admin" and forecast.user_id != user_id:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="You do not have permission to access this forecast"
        )

    return forecast


def build_report_summary(db: Session, forecast: ForecastHistory):
    dataset = db.query(Dataset).filter(Dataset.id == forecast.dataset_id).first()

    predictions = parse_json(forecast.predictions, [])
    comparison_result = parse_json(forecast.comparison_result, None)

    best_prediction = None

    if isinstance(predictions, list) and predictions:
        best_prediction = max(
            predictions,
            key=lambda item: item.get("predicted_value", 0)
        )

    summary = {
        "dataset": {
            "id": dataset.id if dataset else None,
            "name": dataset.name if dataset else None,
            "file_name": dataset.file_name if dataset else None,
            "total_rows": dataset.total_rows if dataset else 0,
            "total_columns": dataset.total_columns if dataset else 0,
            "product_category": dataset.product_category if dataset else None,
            "region": dataset.region if dataset else None,
        },
        "forecast": {
            "id": forecast.id,
            "model_name": forecast.model_name,
            "target_column": forecast.target_column,
            "date_column": forecast.date_column,
            "forecast_periods": forecast.forecast_periods,
            "mae": forecast.mae,
            "rmse": forecast.rmse,
            "mape": forecast.mape,
            "created_at": str(forecast.created_at),
        },
        "insights": {
            "highest_predicted_period": best_prediction,
            "total_predicted_demand": round(
                sum(item.get("predicted_value", 0) for item in predictions),
                2
            ) if isinstance(predictions, list) else 0,
            "average_predicted_demand": round(
                sum(item.get("predicted_value", 0) for item in predictions) / len(predictions),
                2
            ) if isinstance(predictions, list) and predictions else 0,
            "model_comparison_available": comparison_result is not None
        },
        "predictions": predictions,
        "comparison_result": comparison_result
    }

    return summary


def create_summary_report(db: Session, forecast_id: int, user_id: int, role: str):
    forecast = get_forecast_or_404(db, forecast_id, user_id, role)
    summary = build_report_summary(db, forecast)

    report = Report(
        user_id=user_id,
        forecast_id=forecast.id,
        dataset_id=forecast.dataset_id,
        title=f"Forecast Summary Report #{forecast.id}",
        report_type="summary",
        file_path=None,
        summary=json.dumps(summary)
    )

    db.add(report)
    db.commit()
    db.refresh(report)

    create_notification(
        db=db,
        user_id=user_id,
        title="Report Generated",
        message="Forecast summary report generated successfully.",
        type="success"
    )

    return report


def create_excel_report(db: Session, forecast_id: int, user_id: int, role: str):
    os.makedirs(REPORT_DIR, exist_ok=True)

    forecast = get_forecast_or_404(db, forecast_id, user_id, role)
    summary = build_report_summary(db, forecast)

    file_name = f"forecast_report_{forecast.id}_{datetime.now().strftime('%Y%m%d%H%M%S')}.xlsx"
    file_path = os.path.join(REPORT_DIR, file_name)

    wb = Workbook()

    ws = wb.active
    ws.title = "Summary"

    ws.append(["Advanced AI Demand Forecasting Report"])
    ws.append([])
    ws.append(["Forecast ID", forecast.id])
    ws.append(["Dataset Name", summary["dataset"]["name"]])
    ws.append(["Model Name", forecast.model_name])
    ws.append(["Target Column", forecast.target_column])
    ws.append(["Forecast Periods", forecast.forecast_periods])
    ws.append(["MAE", forecast.mae])
    ws.append(["RMSE", forecast.rmse])
    ws.append(["MAPE", forecast.mape])
    ws.append(["Generated At", str(datetime.now())])

    pred_ws = wb.create_sheet("Predictions")
    pred_ws.append(["Period", "Predicted Value"])

    for item in summary["predictions"]:
        pred_ws.append([
            item.get("period"),
            item.get("predicted_value")
        ])

    if summary["comparison_result"]:
        comp_ws = wb.create_sheet("Model Comparison")
        comp_ws.append(["Model", "MAE", "RMSE", "MAPE"])

        for model in summary["comparison_result"].get("models", []):
            comp_ws.append([
                model.get("model_name"),
                model.get("mae"),
                model.get("rmse"),
                model.get("mape")
            ])

    wb.save(file_path)

    report = Report(
        user_id=user_id,
        forecast_id=forecast.id,
        dataset_id=forecast.dataset_id,
        title=f"Excel Forecast Report #{forecast.id}",
        report_type="excel",
        file_path=file_path,
        summary=json.dumps(summary)
    )

    db.add(report)
    db.commit()
    db.refresh(report)

    create_notification(
        db=db,
        user_id=user_id,
        title="Report Generated",
        message="Excel report generated successfully.",
        type="success"
    )

    return report


def create_pdf_report(db: Session, forecast_id: int, user_id: int, role: str):
    os.makedirs(REPORT_DIR, exist_ok=True)

    forecast = get_forecast_or_404(db, forecast_id, user_id, role)
    summary = build_report_summary(db, forecast)

    file_name = f"forecast_report_{forecast.id}_{datetime.now().strftime('%Y%m%d%H%M%S')}.pdf"
    file_path = os.path.join(REPORT_DIR, file_name)

    doc = SimpleDocTemplate(file_path, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph("Advanced AI Demand Forecasting Report", styles["Title"]))
    elements.append(Spacer(1, 12))

    details = [
        ["Forecast ID", forecast.id],
        ["Dataset", summary["dataset"]["name"]],
        ["Model", forecast.model_name],
        ["Target Column", forecast.target_column],
        ["Forecast Periods", forecast.forecast_periods],
        ["MAE", forecast.mae],
        ["RMSE", forecast.rmse],
        ["MAPE", forecast.mape],
        ["Generated At", str(datetime.now())],
    ]

    table = Table(details)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))

    elements.append(table)
    elements.append(Spacer(1, 16))

    elements.append(Paragraph("Forecast Predictions", styles["Heading2"]))

    prediction_data = [["Period", "Predicted Value"]]

    for item in summary["predictions"]:
        prediction_data.append([
            item.get("period"),
            item.get("predicted_value")
        ])

    pred_table = Table(prediction_data)
    pred_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("GRID", (0, 0), (-1, -1), 1, colors.black),
        ("PADDING", (0, 0), (-1, -1), 6),
    ]))

    elements.append(pred_table)

    if summary["comparison_result"]:
        elements.append(Spacer(1, 16))
        elements.append(Paragraph("Model Comparison", styles["Heading2"]))

        comparison_data = [["Model", "MAE", "RMSE", "MAPE"]]

        for model in summary["comparison_result"].get("models", []):
            comparison_data.append([
                model.get("model_name"),
                model.get("mae"),
                model.get("rmse"),
                model.get("mape")
            ])

        comparison_table = Table(comparison_data)
        comparison_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
            ("GRID", (0, 0), (-1, -1), 1, colors.black),
            ("PADDING", (0, 0), (-1, -1), 6),
        ]))

        elements.append(comparison_table)

    doc.build(elements)

    report = Report(
        user_id=user_id,
        forecast_id=forecast.id,
        dataset_id=forecast.dataset_id,
        title=f"PDF Forecast Report #{forecast.id}",
        report_type="pdf",
        file_path=file_path,
        summary=json.dumps(summary)
    )

    db.add(report)
    db.commit()
    db.refresh(report)

    create_notification(
        db=db,
        user_id=user_id,
        title="Report Generated",
        message="PDF report generated successfully.",
        type="success"
    )

    return report


def generate_report(db: Session, forecast_id: int, report_type: str, user_id: int, role: str):
    if report_type == "summary":
        return create_summary_report(db, forecast_id, user_id, role)

    if report_type == "excel":
        return create_excel_report(db, forecast_id, user_id, role)

    if report_type == "pdf":
        return create_pdf_report(db, forecast_id, user_id, role)

    raise HTTPException(
        status_code=status.HTTP_400_BAD_REQUEST,
        detail="Invalid report type. Use summary, excel, or pdf"
    )


def format_report(report: Report):
    return {
        "id": report.id,
        "user_id": report.user_id,
        "forecast_id": report.forecast_id,
        "dataset_id": report.dataset_id,
        "title": report.title,
        "report_type": report.report_type,
        "file_path": report.file_path,
        "summary": parse_json(report.summary, None),
        "created_at": report.created_at
    }